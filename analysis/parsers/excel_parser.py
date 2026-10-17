"""
Excel Multi-Sheet Workbook Parser for Cycler Exports (Landt, Neware, Arbin).
Zero-dependency XML fallback ensures robust parsing even without openpyxl installed.
"""

import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Any, Union

from models import CyclerRawData
from analysis.parsers.base import (
    BaseCyclerParser,
    match_canonical_channel,
    clean_header_str,
    clean_numeric_cell,
    detect_channel_unit,
    is_unit_row,
    parse_duration_to_seconds,
    resolve_best_column_mapping,
)


class ExcelCyclerParser(BaseCyclerParser):
    """Excel workbook parser supporting both openpyxl and zero-dependency fallback."""

    SUPPORTED_EXTENSIONS = frozenset({'.xlsx', '.xlsm', '.xltx', '.xltm'})

    def can_parse(self, filepath: Union[str, Path]) -> bool:
        path = Path(filepath)
        if path.suffix.lower() not in self.SUPPORTED_EXTENSIONS:
            return False
        return zipfile.is_zipfile(path)

    def parse(self, filepath: Union[str, Path]) -> CyclerRawData:
        path = Path(filepath)
        if not path.is_file():
            raise FileNotFoundError(f"File not found: {path}")

        try:
            return self._parse_with_openpyxl(path)
        except Exception:
            return self._parse_with_builtin_zip(path)

    def _parse_with_openpyxl(self, path: Path) -> CyclerRawData:
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet_name = self._select_best_sheet(wb.sheetnames)
        sheet = wb[sheet_name]

        rows = []
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows.append(list(row))
            if len(rows) > 50000:
                break

        wb.close()
        return self._extract_data_from_rows(rows, source_name=path.name, sheet_name=sheet_name)

    def _parse_with_builtin_zip(self, path: Path) -> CyclerRawData:
        """Zero-dependency parser for .xlsx workbooks using built-in zipfile and XML."""
        with zipfile.ZipFile(path, 'r') as zf:
            shared_strings = []
            if 'xl/sharedStrings.xml' in zf.namelist():
                xml_content = zf.read('xl/sharedStrings.xml')
                root = ET.fromstring(xml_content)
                ns = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                for si in root.findall('.//main:si', ns):
                    texts = [t.text or '' for t in si.findall('.//main:t', ns)]
                    shared_strings.append(''.join(texts))

            # 2. Map sheet names to worksheet XML files via workbook.xml and workbook.xml.rels
            sheet_map: Dict[str, str] = {}
            if 'xl/workbook.xml' in zf.namelist() and 'xl/_rels/workbook.xml.rels' in zf.namelist():
                try:
                    wb_root = ET.fromstring(zf.read('xl/workbook.xml'))
                    rels_root = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))

                    rid_to_target: Dict[str, str] = {}
                    for rel in rels_root:
                        rid = rel.attrib.get('Id')
                        target = rel.attrib.get('Target', '')
                        if target:
                            if not target.startswith('xl/'):
                                if target.startswith('/xl/'):
                                    target = target.lstrip('/')
                                else:
                                    target = 'xl/' + target.lstrip('/')
                            rid_to_target[rid] = target

                    for s in wb_root.iter():
                        if s.tag.endswith('sheet'):
                            sname = s.attrib.get('name', '')
                            rid = (
                                s.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                                or s.attrib.get('r:id')
                            )
                            if not rid:
                                for k, v in s.attrib.items():
                                    if (k.endswith('id') or 'id' in k.lower()) and v in rid_to_target:
                                        rid = v
                                        break
                            if sname and rid and rid in rid_to_target:
                                sheet_map[sname] = rid_to_target[rid]
                except Exception:
                    pass

            all_sheet_files = [f for f in zf.namelist() if f.startswith('xl/worksheets/sheet') and f.endswith('.xml')]
            if not all_sheet_files and not sheet_map:
                raise ValueError(f"No worksheet XML found in {path.name}")

            # Prioritize candidate sheets using priority heuristics
            priorities = ['record', 'channel', 'data', 'cycle', 'raw', 'sheet1']
            candidate_pairs: List[Tuple[str, str]] = []
            if sheet_map:
                sorted_names = sorted(
                    sheet_map.keys(),
                    key=lambda name: (
                        min([i for i, p in enumerate(priorities) if p in name.lower()] or [99]),
                        name.lower()
                    )
                )
                for sname in sorted_names:
                    candidate_pairs.append((sname, sheet_map[sname]))
            else:
                for sf in all_sheet_files:
                    candidate_pairs.append((Path(sf).stem, sf))

            # Include any other worksheet files as secondary fallbacks
            seen_files = {p[1] for p in candidate_pairs}
            for sf in all_sheet_files:
                if sf not in seen_files:
                    candidate_pairs.append((Path(sf).stem, sf))

            last_error = None
            for sheet_name, target_sheet_file in candidate_pairs:
                if target_sheet_file not in zf.namelist():
                    continue
                try:
                    sheet_content = zf.read(target_sheet_file)
                    sheet_root = ET.fromstring(sheet_content)
                    ns = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

                    rows_data: List[List[Any]] = []
                    for row_el in sheet_root.findall('.//main:row', ns):
                        current_row: Dict[int, Any] = {}
                        max_col = 0
                        for c in row_el.findall('main:c', ns):
                            r_ref = c.attrib.get('r', '')
                            c_idx = self._col_ref_to_index(r_ref)
                            t_type = c.attrib.get('t', '')
                            val = None

                            if t_type == 'inlineStr':
                                is_el = c.find('main:is/main:t', ns)
                                if is_el is not None and is_el.text is not None:
                                    val = is_el.text
                            else:
                                v_el = c.find('main:v', ns)
                                if v_el is not None and v_el.text is not None:
                                    val_str = v_el.text
                                    if t_type == 's' and shared_strings:
                                        try:
                                            val = shared_strings[int(val_str)]
                                        except (ValueError, IndexError):
                                            val = val_str
                                    elif t_type == 'b':
                                        val = bool(int(val_str))
                                    elif t_type == 'e':
                                        val = None
                                    else:
                                        try:
                                            val = float(val_str)
                                        except ValueError:
                                            val = val_str

                            current_row[c_idx] = val
                            if c_idx > max_col:
                                max_col = c_idx

                        row_list = [current_row.get(i) for i in range(max_col + 1)]
                        if any(x is not None for x in row_list):
                            rows_data.append(row_list)
                        if len(rows_data) > 50000:
                            break

                    return self._extract_data_from_rows(rows_data, source_name=path.name, sheet_name=sheet_name)
                except Exception as e:
                    last_error = e
                    continue

            if last_error:
                raise last_error
            raise ValueError(f"Could not identify valid tabular cycler columns in Excel workbook {path.name}")

    def _col_ref_to_index(self, ref: str) -> int:
        match = re.match(r'([A-Z]+)', ref)
        if not match:
            return 0
        letters = match.group(1)
        idx = 0
        for char in letters:
            idx = idx * 26 + (ord(char) - ord('A') + 1)
        return idx - 1

    def _select_best_sheet(self, sheetnames: List[str]) -> str:
        priorities = ['record', 'channel', 'data', 'cycle', 'raw', 'sheet1']
        for p in priorities:
            for s in sheetnames:
                if p in s.lower():
                    return s
        return sheetnames[0] if sheetnames else "Sheet1"

    def _extract_data_from_rows(self, rows: List[List[Any]], source_name: str, sheet_name: str) -> CyclerRawData:
        header_idx = None
        data_start_idx = 1
        col_mapping: Dict[int, str] = {}
        header_units: Dict[int, str] = {}
        raw_headers: List[str] = []

        for r_idx, row in enumerate(rows[:30]):
            matches: Dict[int, str] = {}
            units: Dict[int, str] = {}

            # Check if next row is a unit row
            next_row = rows[r_idx + 1] if r_idx + 1 < len(rows) else []
            next_is_unit = is_unit_row(next_row)

            for c_idx, cell in enumerate(row):
                if cell is None:
                    continue
                c_str = str(cell).strip()
                canonical = match_canonical_channel(c_str)

                # Try composite header if single cell didn't match
                if not canonical and c_idx < len(next_row) and next_row[c_idx] is not None:
                    comp_str = f"{c_str} {str(next_row[c_idx]).strip()}"
                    canonical = match_canonical_channel(comp_str)

                if not canonical:
                    continue

                matches[c_idx] = canonical
                unit_str = str(next_row[c_idx]) if c_idx < len(next_row) and next_row[c_idx] is not None else None
                detected_u = detect_channel_unit(c_str, unit_str)
                if detected_u:
                    units[c_idx] = detected_u

            if len(matches) >= 2:
                header_idx = r_idx
                raw_headers = [str(c) if c is not None else '' for c in row]
                col_mapping = resolve_best_column_mapping(matches, raw_headers)
                header_units = {k: units[k] for k in col_mapping if k in units}
                data_start_idx = r_idx + 2 if next_is_unit else r_idx + 1
                break

        if header_idx is None:
            raise ValueError(f"Could not identify valid tabular cycler columns in Excel workbook {source_name}")

        time_s: List[float] = []
        voltage_v: List[float] = []
        current_ma: List[float] = []
        step_index: List[int] = []
        cycle_number: List[int] = []
        cap_charge: List[float] = []
        cap_discharge: List[float] = []
        coulombic_efficiency: List[float] = []
        frequency_hz: List[float] = []
        z_real_ohm: List[float] = []
        z_imag_ohm: List[float] = []

        has_step = any(v == 'step_index' for v in col_mapping.values())
        has_cycle = any(v == 'cycle_number' for v in col_mapping.values())
        has_cap_chg = any(v == 'capacity_charge' for v in col_mapping.values())
        has_cap_dchg = any(v == 'capacity_discharge' for v in col_mapping.values())
        has_generic_cap = any(v == 'capacity' for v in col_mapping.values())
        has_ce = any(v == 'coulombic_efficiency' for v in col_mapping.values())
        has_voltage = any(v == 'voltage' for v in col_mapping.values())
        has_freq = any(v == 'frequency' for v in col_mapping.values())
        has_zreal = any(v == 'z_real' for v in col_mapping.values())
        has_zimag = any(v == 'z_imag' for v in col_mapping.values())

        is_cycle_summary = (
            not has_voltage
            and not has_zreal
            and (has_cycle or has_ce)
            and (has_generic_cap or has_cap_dchg or has_cap_chg or has_ce)
        )

        metadata = {
            'source_file': source_name,
            'sheet_name': sheet_name,
            'header_row': header_idx,
            'data_start_row': data_start_idx,
            'raw_columns': raw_headers,
            'mapped_channels': {raw_headers[k]: v for k, v in col_mapping.items() if k < len(raw_headers)},
            'detected_units': {raw_headers[k]: header_units[k] for k in header_units if k < len(raw_headers)},
            'is_cycle_summary': is_cycle_summary,
        }

        for row in rows[data_start_idx:]:
            row_vals: Dict[str, float] = {}
            try:
                for col_idx, canon_name in col_mapping.items():
                    val_raw = row[col_idx] if col_idx < len(row) else None
                    if val_raw is None:
                        continue
                    unit = header_units.get(col_idx)

                    if canon_name == 'time':
                        val = parse_duration_to_seconds(val_raw, unit=unit)
                        if val is None or val == 0.0:
                            val = clean_numeric_cell(val_raw)
                    elif canon_name == 'step_time':
                        continue
                    else:
                        val = clean_numeric_cell(val_raw)
                        if val is not None:
                            if canon_name == 'current':
                                if unit == 'a':
                                    val *= 1000.0
                                elif unit == 'ua':
                                    val /= 1000.0
                            elif canon_name in ('capacity_charge', 'capacity_discharge', 'capacity'):
                                if unit == 'ah':
                                    val *= 1000.0
                                elif unit == 'uah':
                                    val /= 1000.0
                            elif canon_name == 'voltage':
                                if unit == 'mv':
                                    val /= 1000.0

                    if val is not None:
                        row_vals[canon_name] = val

                if not row_vals:
                    continue

                if is_cycle_summary:
                    cyc = int(row_vals.get('cycle_number', len(cycle_number) + 1))
                    cap_val = float(row_vals.get('capacity_discharge', row_vals.get('capacity', row_vals.get('capacity_charge', 0.0))))
                    ce_val = float(row_vals.get('coulombic_efficiency', 99.5)) if has_ce else 99.5
                    if 0.0 < ce_val <= 1.05:
                        ce_val *= 100.0

                    chg_val = float(row_vals.get('capacity_charge', 0.0))
                    if not has_cap_chg and ce_val > 0:
                        chg_val = round(cap_val / (ce_val / 100.0), 4)

                    time_s.append(float((cyc - 1) * 3600.0))
                    voltage_v.append(1.2)  # Nominal full-cell potential for schema integrity
                    current_ma.append(1.0)
                    cycle_number.append(cyc)
                    cap_discharge.append(cap_val)
                    cap_charge.append(chg_val)
                    if has_ce or ce_val is not None:
                        coulombic_efficiency.append(ce_val)
                    continue

                if 'voltage' not in row_vals and 'z_real' not in row_vals:
                    continue

                curr_val = row_vals.get('current', 0.0)
                volt_val = row_vals.get('voltage', 0.0)
                time_s.append(row_vals.get('time', float(len(voltage_v))))
                voltage_v.append(volt_val)
                current_ma.append(curr_val)

                if has_step:
                    step_index.append(int(row_vals.get('step_index', 0)))
                if has_cycle:
                    cycle_number.append(int(row_vals.get('cycle_number', 0)))

                if has_cap_chg:
                    cap_charge.append(float(row_vals.get('capacity_charge', 0.0)))
                if has_cap_dchg:
                    cap_discharge.append(float(row_vals.get('capacity_discharge', 0.0)))
                if has_generic_cap and not has_cap_chg and not has_cap_dchg:
                    # Route single capacity column based on current sign
                    q_val = float(row_vals.get('capacity', 0.0))
                    if curr_val < -0.005:
                        cap_discharge.append(q_val)
                        cap_charge.append(0.0)
                    else:
                        cap_charge.append(q_val)
                        cap_discharge.append(0.0)

                if has_ce and 'coulombic_efficiency' in row_vals:
                    ce_v = float(row_vals['coulombic_efficiency'])
                    if 0.0 < ce_v <= 1.05:
                        ce_v *= 100.0
                    coulombic_efficiency.append(ce_v)

                if has_freq:
                    frequency_hz.append(float(row_vals.get('frequency', 0.0)))
                if has_zreal:
                    z_real_ohm.append(float(row_vals.get('z_real', 0.0)))
                if has_zimag:
                    z_imag_ohm.append(float(row_vals.get('z_imag', 0.0)))
            except (ValueError, TypeError, IndexError):
                continue

        if not time_s:
            raise ValueError(f"No valid numeric data points parsed from Excel sheet {sheet_name}")

        if time_s and time_s[0] > 1e7:
            t0 = time_s[0]
            time_s = [round(t - t0, 4) for t in time_s]

        final_has_cap_chg = has_cap_chg or (has_generic_cap and len(cap_charge) == len(time_s)) or (is_cycle_summary and len(cap_charge) == len(time_s))
        final_has_cap_dchg = has_cap_dchg or (has_generic_cap and len(cap_discharge) == len(time_s)) or (is_cycle_summary and len(cap_discharge) == len(time_s))
        final_has_ce = len(coulombic_efficiency) == len(time_s) and len(coulombic_efficiency) > 0

        if is_cycle_summary:
            metadata['summary_coulombic_efficiency'] = coulombic_efficiency
            metadata['summary_cycle_number'] = cycle_number
            metadata['summary_capacity'] = cap_discharge

        return CyclerRawData(
            time_s=time_s,
            voltage_v=voltage_v,
            current_ma=current_ma,
            step_index=step_index if has_step else None,
            cycle_number=cycle_number if (has_cycle or is_cycle_summary) else None,
            capacity_charge_mah=cap_charge if final_has_cap_chg else None,
            capacity_discharge_mah=cap_discharge if final_has_cap_dchg else None,
            coulombic_efficiency=coulombic_efficiency if final_has_ce else None,
            frequency_hz=frequency_hz if has_freq else None,
            z_real_ohm=z_real_ohm if has_zreal else None,
            z_imag_ohm=z_imag_ohm if has_zimag else None,
            metadata=metadata,
        )
